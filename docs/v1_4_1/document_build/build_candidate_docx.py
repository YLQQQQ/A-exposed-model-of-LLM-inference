"""将 ExposedPath 中文 Markdown 候选稿构建为可审阅的 DOCX。"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ACCENT = "1F4E78"
LIGHT = "D9EAF7"
GRID = "AAB7C4"
TEXT = RGBColor(31, 31, 31)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def set_run_font(run, east_asia: str, latin: str, size: float, bold: bool | None = None) -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = TEXT
    if bold is not None:
        run.bold = bold


def append_inline(paragraph, text: str, size: float = 10.5) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "等线", "Consolas", size - 0.5)
            run.font.color.rgb = RGBColor(88, 58, 16)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, "黑体", "Arial", size, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, "宋体", "Times New Roman", size)


def configure_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in (
        ("Title", 24, ACCENT),
        ("Heading 1", 16, ACCENT),
        ("Heading 2", 13, "2F5597"),
        ("Heading 3", 11.5, "365F91"),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    styles["Title"].paragraph_format.space_before = Pt(70)
    styles["Title"].paragraph_format.space_after = Pt(24)
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    core = doc.core_properties
    core.title = title
    core.subject = "ExposedPath 研究文档候选版"
    core.author = "ExposedPath 项目组"
    core.comments = "由仓库内 Markdown 权威候选源生成；尚未完成 Protocol Freeze。"

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("ExposedPath · Pre-Pilot 候选版")
    set_run_font(run, "等线", "Arial", 8)
    run.font.color.rgb = RGBColor(117, 117, 117)
    add_page_number(section.footer.paragraphs[0])


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(cells)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if row_index == 0:
                set_cell_shading(cell, ACCENT)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F4F8FB")
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            value = values[col_index] if col_index < len(values) else ""
            append_inline(paragraph, value, 8.5)
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build(markdown_path: Path, output_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    title = next((line[2:].strip() for line in lines if line.startswith("# ")), markdown_path.stem)
    doc = Document()
    configure_document(doc, title)

    index = 0
    cover_open = True
    while index < len(lines):
        raw = lines[index].rstrip()
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("| "):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                paragraph = doc.add_paragraph(style="Title")
                append_inline(paragraph, text, 24)
            else:
                if cover_open:
                    doc.add_page_break()
                    cover_open = False
                paragraph = doc.add_paragraph(style=f"Heading {level - 1}")
                append_inline(paragraph, text, 16 if level == 2 else 13)
            index += 1
            continue
        if line.startswith("**") and line.endswith("**"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            append_inline(paragraph, line, 11)
            index += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if bullet:
            paragraph = doc.add_paragraph(style="List Bullet")
            append_inline(paragraph, bullet.group(1))
        elif numbered:
            paragraph = doc.add_paragraph(style="List Number")
            append_inline(paragraph, numbered.group(2))
        elif line.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.right_indent = Cm(0.8)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(8)
            append_inline(paragraph, line[2:], 11)
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor.from_string(ACCENT)
        else:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(0.74)
            append_inline(paragraph, line)
        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("用法：build_candidate_docx.py 输入.md 输出.docx")
    build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())


if __name__ == "__main__":
    main()
