"""基于 v6.0/v2.0 母版生成当前 ExposedPath 研究设计与实验协议。"""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt, RGBColor


DESIGN_SHA256 = "0EE058EE1DD2D1FB69EC59D230DBD6B6F03BAEDC319D9C246AEE3ED229B530F1"
PROTOCOL_SHA256 = "1AD97E338F03C951E46BD81B15A5633D46A3B6E6490FEC4B2DFA8840028C0685"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_paragraph(doc: Document, index: int, text: str) -> None:
    set_paragraph_text(doc.paragraphs[index - 1], text)


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    result = Paragraph(new_p, paragraph._parent)
    if style:
        result.style = style
    result.add_run(text)
    return result


def insert_block_after(paragraph: Paragraph, entries: list[tuple[str, str | None]]) -> None:
    cursor = paragraph
    for text, style in entries:
        cursor = insert_after(cursor, text, style)


def find_paragraph(doc: Document, exact_text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise ValueError(f"未找到段落：{exact_text}")


def replace_in_all_text(doc: Document, old: str, new: str) -> int:
    count = 0
    for paragraph in doc.paragraphs:
        if old in paragraph.text:
            set_paragraph_text(paragraph, paragraph.text.replace(old, new))
            count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if old in paragraph.text:
                        set_paragraph_text(paragraph, paragraph.text.replace(old, new))
                        count += 1
    return count


def replace_in_headers_and_footers(doc: Document, old: str, new: str) -> int:
    """替换页眉页脚中的版本文本，并避免重复处理链接到前节的部件。"""
    count = 0
    seen_parts: set[int] = set()
    for section in doc.sections:
        # 母版未启用首页/奇偶页专用页眉页脚；只访问已使用的默认部件，
        # 避免 python-docx 因读取不存在的可选部件而新建空关系。
        containers = (section.header, section.footer)
        for container in containers:
            part_key = id(container.part)
            if part_key in seen_parts:
                continue
            seen_parts.add(part_key)
            for paragraph in container.paragraphs:
                if old in paragraph.text:
                    set_paragraph_text(paragraph, paragraph.text.replace(old, new))
                    count += 1
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            if old in paragraph.text:
                                set_paragraph_text(paragraph, paragraph.text.replace(old, new))
                                count += 1
    return count


def ensure_table_rows(table, count: int) -> None:
    while len(table.rows) < count:
        table.add_row()


def set_table(table, rows: list[list[str]]) -> None:
    ensure_table_rows(table, len(rows))
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value


def format_background_comparison_table(table) -> None:
    """压缩第一章比较表，并清除母版遗留的黄色修订底色。"""
    for row_index, row in enumerate(table.rows):
        fill = "D9EAF7" if row_index == 0 else ("FFFFFF" if row_index % 2 else "F5F8FA")
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            shading = tc_pr.find(qn("w:shd"))
            if shading is None:
                shading = OxmlElement("w:shd")
                tc_pr.append(shading)
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:fill"), fill)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    run.bold = row_index == 0


def remove_title_rule(doc: Document) -> None:
    """移除母版 Title 样式自带的下边框，保留纯标题排版。"""
    title_style = doc.styles["Title"]
    paragraph_properties = title_style.element.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def keep_table_row_together(row) -> None:
    """避免表格行在分页处只留下极少文本。"""
    row_properties = row._tr.get_or_add_trPr()
    if row_properties.find(qn("w:cantSplit")) is None:
        row_properties.append(OxmlElement("w:cantSplit"))


def normalize_revision_marks(doc: Document) -> None:
    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        if style_name in doc.styles:
            doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.highlight_color = None
            run.font.color.rgb = RGBColor(0, 0, 0)
    blue_marks = {"0070C0", "123E62", "163758", "3C5569", "5A5F64"}
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.highlight_color = None
                        color = run.font.color.rgb
                        if color is not None and str(color).upper() in blue_marks:
                            run.font.color.rgb = RGBColor(0, 0, 0)
    # 文本框或其他 OOXML 容器中的 run 不一定出现在 python-docx 的
    # paragraphs/tables 视图中，仍需在文档树上统一关闭可见高亮。
    for highlight in doc.element.xpath(".//w:highlight"):
        highlight.set(qn("w:val"), "none")


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def patch_design(source: Path, output: Path) -> None:
    if sha256(source) != DESIGN_SHA256:
        raise RuntimeError("研究设计母版哈希不匹配，拒绝继续")
    shutil.copy2(source, output)
    doc = Document(output)

    replacements = {
        2: "LLM 推理请求可见暴露研究设计",
        3: "当前研究设计主体",
        4: "研究对象：单 GPU、请求内部、纯模型推理的 Host–accelerator request-visible exposure",
        6: "版本定位：本文件已经逐章归并 v1.4.1 的研究背景、研究立意、测量语义、信息增量、决策增量与执行路线，作为当前研究设计主体。方法链固定为 Raw→S→{A,B}→D/Exposure Signature。具体执行合同由《ExposedPath 实验协议》承接；当前代码和历史 trace 只表示 Prototype/Engineering 事实。",
        7: "版本：v7.1　日期：2026 年 9 月 10 日　状态：当前研究设计主体",
        8: "v1.4.1 整合结论：Activity Cost 不等于 Request-Visible Exposure；S 层必须依据 synchronization/completion semantics 恢复完整语义前驱集合 W(s)、terminal 证据与 validity，时间重叠本身不能证明 dependency；研究证据按 Correctness→Information Gain→Decision Gain 逐级建立。",
        9: "平台边界：当前方法以单 GPU、请求内部 Host–device execution domain 为核心；RTX 4090、RTX 6000 Ada 及其他 GPU 均须先通过正式平台资格检查，不能在 Pre-Pilot 阶段写成已冻结主平台。",
        11: "贡献一：方法候选。提出 ExposedPath，通过同步完成语义把 Raw trace 转换为可审计的 W(s)、terminal 与 validity，并在此基础上形成 request/phase 互斥墙钟 accounting 和 per-sync provenance；该贡献只有在 Q0 通过后才能表述为已验证方法。",
        12: "贡献二：信息增量候选。通过 N1 与 G1 检验传统 activity 指标与 request-visible exposure 在哪些条件下近似等价、稳定非等价或依赖执行区间变化；结果可以是比例、非比例、条件性或 null，主张随证据收缩。",
        13: "贡献三：决策增量候选。通过 G2 的预注册真实优化与 held-out 比较，检验 Exposure Signature 是否相对传统指标增加优化解释或优先级信息；若无增量，只保留正确性与信息边界，不声称优化预测或决策优势。",
        26: "1.1 研究背景：用户可见时延来自异步 Host–Accelerator 协作",
        27: "交互式 LLM 推理的用户感受首先体现在 request latency、Prefill latency、Decode latency 和逐 Token 生成速度。但一次模型推理并不是一串彼此孤立的 GPU kernel。Host 侧的框架与控制路径持续推进程序逻辑并提交 CUDA 工作，GPU 异步执行 kernel 和 MemOp，两侧可以重叠推进，最后在 synchronization/completion boundary 上汇合。因此，纯模型推理本身就是时延敏感的异步 Host–accelerator 执行系统。理解用户可见时延，需要解释这条协作链如何形成 request/phase wall-clock。",
        28: "Nsight Systems、CUPTI 等工具能够可靠报告 kernel duration、GPU active/utilization、CUDA API duration、raw synchronization duration、API/kernel count 和 launch-to-start 等执行事实。这些事实描述 Activity Cost，即系统做了多少工作；它们不等于 Request-Visible Exposure，即这些异步工作中有多少真正暴露到请求或阶段墙钟。长 kernel 可能在 Host 开始等待前已完成大部分执行，长 sync 也可能只是等待边界发生迁移。ExposedPath 要解决的核心矛盾是 Activity Cost 不等于 Request-Visible Exposure。",
        29: "1.1.1 传统性能指标与 ExposedPath 的测量对象",
        30: "传统 trace 和 profiler 主要回答发生了哪些活动、哪里耗时以及设备是否繁忙。ExposedPath 在这些事实之上增加一层同步语义解释：哪些活动属于某次同步返回前必须完成的集合，它们有多少已经隐藏完成，有多少真正暴露为 request/phase 的设备等待。下表说明两类测量对象的关系。",
        31: "1.2 现有研究脉络与 ExposedPath 的位置",
        32: "已有研究分别解释 GPU 硬件与 kernel 行为、CPU–GPU 提交和排队、跨栈 profiling、LLM 阶段时延以及 serving 优化。这些工作为 ExposedPath 提供 Raw facts、硬件解释和系统背景，但尚未把异步 activity 按 CUDA synchronization/completion semantics 映射为互斥、守恒且可审计的 request-visible exposure。ExposedPath 补的是这一测量层，不替代 Nsight Compute、serving engine 或既有优化方法。",
        33: "1.3 核心学术缺口：Activity 到 Exposure 之间缺少可验证映射",
        34: "缺口 1：Activity Cost 与 Request-Visible Exposure 是不同测量对象。完整 kernel/API/sync duration、利用率或事件数量不能直接解释 request/phase 墙钟。",
        35: "缺口 2：Temporal overlap 不等于 synchronization dependency。Stream、device/context 和 event synchronization 的 completion condition 不同，不能仅凭时间重叠判断谁在等待谁。",
        36: "缺口 3：Raw trace 只提供执行事实，仍需 S 层恢复 W(s)、terminal 和 validity。缺少 identity、ownership、提交顺序或 completion evidence 时，分析必须 fail closed。",
        37: "缺口 4：完整 request/phase accounting 与单次同步 provenance 回答不同问题。A 需要互斥并守恒于墙钟，B 只解释单次同步的 hidden progress、sync-overlap、terminal 和 return tail，不能跨同步直接求和。",
        38: "缺口 5：现有指标是否足够不能靠预设判断。需要在冻结 workload 和 phase 下检验 activity 与 exposure 是近似成比例、稳定非成比例、随执行区间变化，还是基本等价。",
        39: "缺口 6：测量信息只有在改变解释或决策时才形成更强系统价值。需要用真实、预注册的优化干预检验 Exposure Signature 是否相对传统指标增加 held-out 决策信息。",
        40: "缺口 7：正确性、信息增量和决策增量必须分层取证。后层结果不能补偿错误的同步语义，也不能根据正式结果反向修改指标、工作负载或主张。",
        41: "1.4 ExposedPath 的研究立意与贡献边界",
        42: "ExposedPath 不是新的 serving engine，也不是恢复全部 CUDA 因果关系的通用 profiler。它是一种 trace-driven measurement methodology：先根据同步完成语义恢复 W(s)、terminal 与 validity，再用 A 解释 request/phase 墙钟去了哪里，用 B 解释单次同步如何形成，最后以 D 和 Exposure Signature 导航后续机制审计。",
        43: "研究证据按 Correctness→Information Gain→Decision Gain 递进。Q0 用受控 CUDA 微程序和独立 oracle 验证测量语义；N1/G1 分别以人为 completion-boundary 干预和自然 workload 检验传统指标何时充分、何时失配；G2 在冻结规则和 held-out 场景下检验决策增量。研究不预设必须出现强反转或反常识结果；若传统指标在所测范围内已经充分，或 G2 没有增量，应如实收缩相应 claim。",
        50: "A/S/B 的核心 claim 只针对 T_model-request。流式推理中，为确认模型侧 Token 已就绪而发生的自然逐 Token 同步属于 Decode 完成边界；Token 就绪后的反分词、文本拼接、网络传输、前端显示和外部消费者处理不进入核心 A。tokenization、服务排队和外部 I/O 只能作为外层辅助项，不能归入 A_host_path。",
        53: "W/M/P/C 是实验因素，而不是四套互不相容的指标。核心 W 包含 fixed_input_tokens、fixed_output_tokens 与 batch_size；当前不研究请求并发或 continuous batching，但保留纯模型流式推理中自然逐 Token 的 Token-ready completion boundary。所有 invocation 隔离执行。batch=1 可表述为单请求；batch>1 的 E2E 与 A/B/D 对应整次 batched inference invocation，不分摊成单样本因果开销。",
        56: "vLLM、llama.cpp 等完整推理引擎会同时改变调度、KV cache、算子后端、stream 与 graph 组织，只能作为后续完整 computing-stack endpoint，不能替代 G2 的同栈单变量干预。具体候选与删减规则见实验协议 v2.1；在 Q0、N1/G1、G2 主线完成前不启动这些扩展。",
        57: "研究设计固定“同栈、单主干预、配置可审计、结果允许失败”的原则；具体 compile/graph 模式、capture 范围、warmup、repeat 与删除条件由 Pilot 形成证据后写入 Protocol Freeze，不能在 Pre-Pilot 阶段凭经验固定。",
        65: "资源有限时，优先保证 Q0→N1/G1→G2 主证据链。G3/G4 只在核心链完成后用于检验一个必要的适用边界；外部平台、端侧、MoE、量化和跨引擎扩展不得挤占核心工作。",
        68: "G2 是 Decision Gain 层，不是 Q0/G1 的替代品。若 G2 干预不可运行或 held-out 结果无增量，应删除预测或决策优势主张；不能通过更换跨栈 endpoint 或事后改分数追求正结果。",
        69: "当前不启用任何跨栈或事后选择的 G2 fallback。C1、G2_score、held-out 单元和 baseline 在本阶段都只是待 Pilot 审查的候选；只有在 Formal 数据前完成预注册与 Protocol Freeze 后，才能用于决策增量检验。",
        83: "Raw trace 只记录 API、kernel、MemOp、synchronization、identity 与时间事实。S 按同步类型、提交顺序、context/stream/event、ownership 和 observation contract 恢复 W(s)、terminal 证据与 validity。A 与 B 共同依赖 S：A 回答 request/phase 的互斥墙钟去向，B 回答单次同步的来源与隐藏/暴露结构；D 与 Exposure Signature 只能从冻结后的 A/B 派生。",
        93: "W(s) = {e | e 属于同步 s 的 completion scope，且在 s 之前按冻结提交规则成为其语义前驱并可归属到当前请求}",
        95: "S 的最小产物包括 completion_scope、W(s)、terminal evidence、validity、reason code、ownership 和 supported-sync coverage。W(s) 包含同步入口前已经完成但仍属于语义前驱的活动；证据不足、映射冲突或 terminal 不唯一时必须 fail closed。",
        110: "W(s) := {e ∈ Pred_semantic(s) | submission_proven(e,s) ∧ ownership_supported(e,s)}。e.end 与 s.start 的关系用于计算 hidden/exposed 状态，不用于决定 e 是否属于 W(s)。",
        113: "Validity 至少区分 VALID_NONEMPTY、VALID_EMPTY、AMBIGUOUS 与 INVALID。VALID_EMPTY 只表示不存在请求内、受支持且可归属的语义前驱；相关工作已在 sync.start 前完成时，W(s) 仍然非空，其执行进展记为 hidden，而不是合法空集合。",
        114: "最小输出字段必须覆盖 sync identity、completion scope、wait-set identity、terminal evidence、validity/status、ownership 与 reason codes。AMBIGUOUS 与 INVALID 不得静默转成零值；B-valid coverage、有效空集合和拒绝回答的比例分别报告。",
        123: "对证据充分的同步：A_device_wait(s)=μ([s.start,s.end) ∩ ⋃e∈W(s)[e.start,e.end))；A_sync_residual(s)=duration(s)−A_device_wait(s)。μ 表示 interval union 墙钟长度；同步前已完成成员仍保留在 W(s)，但对该 sync 的暴露交集可以为 0。",
        124: "对 VALID_EMPTY，同步确无请求内受支持前驱，A_device_wait=0；对 AMBIGUOUS/INVALID，相关同步区间进入 A_unattributed 或合同指定的未知类别，不能为闭合塞入 A_host_path。具体 residual 规则由 Measurement Contract 冻结。",
        128: "分析器必须输出分区互斥/守恒、B 不进入 A 预算、completed-before-sync 成员保留、无关重叠排除、validity 传播和 invalid-sync 不归 Host 等机器可检查 QA 字段；字段名以 Measurement Contract 的冻结 schema 为准。",
        131: "Reason taxonomy 至少覆盖 trace 完整性、ownership、支持性、mapping、submission/closure 与 terminal 证据。primary/secondary reason 的优先级必须版本化；coverage 同时按 count、sync duration 和 B-valid duration 报告。",
        153: "2.4.2 D 的导航作用与 Exposure Signature",
        158: "D 只用于定位 Host/API 与 device-wait 的主导区域，不能证明 CPU/GPU 硬件根因或可实现 speedup。核心解释必须继续下钻到 A 二级组成、per-sync W(s)/B provenance 和必要的代表 activity 审计；Exposure Signature 只能由冻结后的 A/B 字段派生。",
        166: "terminal 不是简单的全局最大 end_ns。只有同步 completion scope、依赖关系、identity/ownership 和时间容差共同支持唯一完成边界时，才能给出 terminal；并列或证据不足时标记 ambiguous/invalid。关联 launch API 只提供附加 provenance，不决定核心 terminal validity。",
        169: "terminal(s) = unique completion evidence for W(s) under the frozen synchronization semantics",
        170: "在冻结规则允许时，terminal 可以由 W(s) 中最后满足 completion condition 的活动表示，但必须先证明候选属于语义依赖并且在时间容差内唯一。terminal 的完整执行可在 sync 前部分或全部完成，因此 terminal duration 不等于同步阻塞时间。",
        184: "A 的互斥闭合只是 partition implementation invariant。Q0 的核心正确性证据必须来自受控 CUDA 用例中的 W(s)/terminal 对独立 oracle 精确匹配、completed-before-sync 与无关重叠负例、validity/fail-closed 和真实 observation coverage。",
        189: "3.1 当前阶段：Engineering 与 Measurement Contract",
        191: "当前阶段结束时应能回答：每个 Token/phase 的可观察完成边界是什么；各 sync completion scope 如何定义；W(s)、terminal 和 validity 如何恢复；A/B 哪些量可加；缺证据时怎样拒绝回答。目前这些内容尚未全部冻结。",
        192: "3.2 从离线语义设计到 Pilot 与 Protocol Freeze",
        195: "执行顺序为：封存 Prototype→Measurement Contract→Q0 oracle→Canonical Raw→S→A/B→D/Signature→Q0→runner 对齐→Engineering Pilot→平台资格→OOM/可行域→Pilot→Protocol Freeze→N1/G1→G2。Pilot 只用于确定运行政策，Formal 数据必须在 Freeze 后重新采集。",
        196: "3.2.2 与实验协议 v2.1 的职责接口",
        198: "候选 sentinel、warmup/repeat、运行顺序、失败重跑、coverage 门槛、G2 候选和停止规则由实验协议 v2.1 维护，并由平台资格、OOM 探测与 Pilot 提供依据；研究设计只维护研究问题、方法语义和证据—主张边界。",
        199: "3.2.3 Measurement Contract 与阶段 Gate",
        201: "3.3 正式证据主线：Q0→N1/G1→G2",
        203: "Q0 是 correctness 硬门；N1 以人为 completion-boundary intervention 检验局部指标，G1 在自然逐 Token 同步与预定义 workload 下检验 activity→exposure 的比例、非比例或区间依赖关系；G2 才检验 held-out decision gain。G3/G4 为可选适用边界，G5–G7 延后。",
        204: "若 compile/graph 在 Pilot 中输出不一致、graph break/capture 不稳定或无法形成可审计同栈干预，则删除或收缩 G2，不修改框架源码、不临时改用跨栈 endpoint，也不改变 A/B 测量语义。",
        205: "3.4 三幅核心图、论文证据链与 Artifact",
        215: "高水平 systems/measurement 研究不以是否自研算子作为唯一标准。本项目首先要证明测量对象新且正确，其次证明相对常规 activity 指标存在稳定、非冗余的信息，或明确界定常规指标已经充分的执行区间；若这些信息还能改变优化解释或优先级，才形成更强的 Decision Gain。强反转或反常识现象只是可能结果，不是预设成功条件。",
        229: "ExposedPath 面向单 GPU、请求内部的纯模型推理，依据 synchronization/completion semantics 把 activity 映射为 request-visible exposure：S 恢复 W(s)、terminal 证据与 validity；A 形成 request/phase 互斥墙钟 accounting；B 保留 per-sync hidden/exposed provenance；D 负责导航，Exposure Signature 负责从冻结 A/B 汇总机制。论文证据按 Correctness→Information Gain→Decision Gain 逐级成立。",
        233: "程序字段、null/validity 规则、输出 schema 与 legacy migration 由《ExposedPath 实验协议》及 Measurement Contract 维护；本文件保留研究语义和 claim—evidence 边界。",
        261: "附录 D　历史 Prototype 结果与方向性线索（不得作为当前方法证据）",
        262: "本附录只保存冻结前历史 prototype 的双 GPU 输出，用于追溯旧实现和生成回归问题。由于 analyzer、W(s)、terminal、validity、manifest、repeat 和运行合同不符合或无法证明符合 v1.4.1，这些数据不得作为 Q0、Pilot 或 Formal 证据，也不得支持当前方法正确性。",
        264: "历史数据来自同一服务器的 RTX 4090 与 RTX 6000 Ada，各 16 个 Qwen2.5-1.5B workload，prompt={128,256,512,1024}、batch={1,2,4,8}、output=16、每组 5 repeats。它们仅用于 Prototype/Engineering 回归和识别旧实现问题。",
        265: "这些历史输出不能支持正式显著性、平台排名、activity→exposure 规律、当前 A/B 正确性或跨主机因果结论。",
        274: "历史 Prefill 图显示过候选方向，但旧 analyzer 与运行合同尚未满足 v1.4.1，因此该趋势只能生成待检验假设，不能预先冻结为 G1 结论。",
        279: "历史 Decode 图曾显示平台条件差异，但其身份、同步语义和数据资格不足；正式 G1 必须在冻结协议下重新采集，并允许比例、非比例、条件性或 null 结果。",
        281: "旧版 predecessor/gap 与 B 字段不得直接迁移为当前结论。历史 trace 只能用于检查新版 parser/S 的回归行为；只有 Measurement Contract 冻结且 validity 成立的记录，才能报告 terminal hidden/exposed provenance。",
        282: "任何历史 gap 或 terminal 现象都只能作为诊断线索，不能命名 CPU scheduling、driver overhead、GPU queue 或其他根因；若需根因结论，必须增加独立证据。",
        284: "历史 prototype 能生成旧版阶段墙钟和 per-sync 输出，但这不能证明其 W(s)、terminal 或 validity 符合 v1.4.1。",
        285: "历史 Prefill 变化只保留为待 Formal G1 检验的候选假设，不能作为信息增量证据。",
        286: "历史 Decode 平台差异只保留为候选线索，不能写成当前研究发现。",
        287: "历史 trace 可用于回归和选择 Q0 负例，但不能提升为 Pilot/Formal，也不能证明传统指标不足。",
        294: "D.8 当前处置",
        295: "历史产物已封存为 Prototype/Engineering 证据。当前唯一最高优先级是完成 Measurement Contract v0.2，随后依次构造 Q0 oracle、Canonical Raw、S 与 A/B；在 Q0 通过、平台资格和 Pilot 完成前，不采集或解释正式 N1/G1/G2。",
    }
    for index, text in replacements.items():
        set_paragraph(doc, index, text)

    protocol_name = "ExposedPath 实验协议"
    replace_in_all_text(doc, "实验与分析协议冻结版 v2.0", protocol_name)
    replace_in_all_text(doc, "实验协议 v2.0", "实验协议 v2.1")
    replace_in_all_text(doc, "Pre-Pilot Freeze", "Pre-Pilot 语义修订")

    set_table(doc.tables[0], [
        ["常见 trace 或 profiler 观察", "通常回答的问题", "ExposedPath 增加的问题", "严格解释边界"],
        ["kernel duration / hotspot", "哪个 kernel 执行最长、热点在哪里", "该 activity 有多少属于 W(s)，多少真正暴露为设备等待", "完整 kernel duration 不是延迟贡献"],
        ["GPU active / utilization", "GPU 是否繁忙、设备工作量是否增加", "设备工作有多少最终进入 request/phase 的 A_device_wait", "繁忙不等于用户正在等待"],
        ["raw sync duration", "某个同步 API 在 Host 上持续多久", "它等待哪些 W(s)，terminal 是谁，等待是新增还是迁移", "局部 sync duration 不自动等于新增 E2E"],
        ["API/kernel count、launch-to-start", "提交密度、启动或排队现象如何", "这些活动在互斥请求墙钟中占多少，是否形成暴露", "时间重叠和 correlation 不能单独证明 dependency"],
    ])
    format_background_comparison_table(doc.tables[0])
    set_table(doc.tables[5], [
        ["证据问题", "要回答的问题", "核心证据"],
        ["Q0 Correctness", "S 是否按 CUDA completion semantics 正确恢复 W(s)、terminal、validity，并使 A/B 满足语义不变量？", "受控 CUDA 微程序、独立 oracle、合成 fixture、真实 observation stack 与 fail-closed"],
        ["N1 Information Gain", "人为移动 completion boundary 时，局部 raw sync 能否代表 request-level 增量？", "自然基线与干预隔离、Pass0 ΔE2E、A 净变化、B waiting migration/hidden progress"],
        ["G1 Information Gain", "自然 workload 变化下，activity 与 exposure 是同比、稳定非同比、区间依赖还是 null？", "冻结 workload、Prefill 绝对量、Decode per-token、A/B/Signature 与传统 baseline"],
        ["G2 Decision Gain", "Exposure 信息能否在 held-out 场景中增加真实优化选择或解释价值？", "预注册干预、baseline、held-out、ΔA 与失败后 claim 收缩"],
    ])
    set_table(doc.tables[6], [
        ["证据层级", "必须具备的证据", "允许主张"],
        ["L1 Correctness", "Q0 全部必需 case 通过", "在声明支持的 observation 条件下正确恢复测量对象"],
        ["L2 Information Gain", "L1 + N1/G1 的稳定非冗余信息或明确等价边界", "相对传统 activity 指标增加解释信息，或界定其充分区间"],
        ["L3 Decision Gain", "L2 + 预注册 held-out G2", "在限定优化与平台范围内增加决策或响应解释价值"],
        ["L4 Applicability", "L1–L3 + 必要的 G3/G4 边界实验", "限定条件下的适用范围，不自动推广到多 GPU/serving"],
    ])
    set_table(doc.tables[7], [
        ["sync 类型", "completion scope / W(s) 规则", "失败条件"],
        ["Stream sync", "目标 stream 在同步前按序提交的请求内语义前驱，包括 sync 入口前已完成者及可观察传递前驱", "缺 stream/context、提交顺序、default-stream 或 closure 证据"],
        ["Device/context sync", "owned device/context 中同步前提交且属于请求范围的语义前驱，包括已经完成的成员", "context/ownership 不唯一或存在无法排除的外部 workload"],
        ["Event sync", "event record boundary 所代表的提交前缀及可观察传递前驱，不含 record 后活动", "event record、event ID 或 wait mapping 不可恢复"],
        ["未知或隐式 sync", "只有 registry 明确其 completion semantics 与必需 identity 时才构造", "否则 unsupported/ambiguous/invalid 并 fail closed"],
    ])
    doc.tables[9].cell(2, 3).text = "被测 owned context 中同步前已证明 submitted 的请求内语义前驱；成员是否在 sync.start 前完成只影响 hidden/exposed，不影响 W(s) 成员资格"
    set_table(doc.tables[10], [
        ["同步类别", "Pred_semantic(s) 的规则", "不能恢复时"],
        ["Stream sync", "目标 stream 在同步前进入排序域的全部请求内工作，以及通过可观察 event/default-stream 关系依赖的传递前驱", "DEPENDENCY_CLOSURE_UNOBSERVABLE"],
        ["Device/context sync", "owned context 内同步前已证明 submitted 的请求内工作，不以 e.end>s.start 过滤", "MULTI_CONTEXT_OWNERSHIP_AMBIGUOUS 或缺失 context"],
        ["Event sync", "event record 节点所代表的提交前缀及其可观察前驱；record 后活动不属于该 event 条件", "MISSING_EVENT_RECORD / MISSING_EVENT_ID"],
    ])
    set_table(doc.tables[11], [
        ["wait_set_status", "W(s)/terminal", "A 处理", "B 处理"],
        ["VALID_NONEMPTY", "W(s) 非空；terminal 只有在 completion evidence 唯一时成立", "用 W(s) 与 sync window 的 interval union 计算 A_device_wait；其余按冻结 residual 规则", "terminal 唯一则 B_VALID，否则转 AMBIGUOUS/INVALID"],
        ["VALID_EMPTY", "没有请求内受支持且可归属的语义前驱；terminal N/A", "A_device_wait=0；其余按冻结 residual 规则", "B_NOT_APPLICABLE，不伪造 terminal"],
        ["AMBIGUOUS", "存在多个合理 completion/ownership/terminal 解释", "不转成零值或已知 device wait；进入未知/不归属处理", "B_AMBIGUOUS，保留候选与原因"],
        ["INVALID", "必需事实缺失、冲突、dropped 或不满足合同", "相关区间进入 A_unattributed 或被排除并记录", "B_INVALID，保留原因"],
    ])
    set_table(doc.tables[22], [
        ["内容", "当前任务", "完成标准", "状态", "产物"],
        ["Measurement Contract", "冻结 Token/phase、sync scope、W(s)、terminal、validity、A/B/D/Signature", "规则可测试、无未决语义", "进行中；Gate 1=FAIL", "Measurement Contract v0.2"],
        ["Q0 oracle", "为正例、负例、含糊例预写独立标准答案", "不调用或复制 analyzer", "未开始", "Q0 case registry + expected"],
        ["离线实现链", "Canonical Raw→S→A/B→D/Signature", "合成 fixture 与不变量测试通过", "未完成", "版本化 schema、代码与测试"],
    ])
    for table_index in (23, 24, 25, 26):
        for row in doc.tables[table_index].rows:
            for cell in row.cells:
                cell.text = cell.text.replace("实验协议 v2.0", "实验协议 v2.1")
    doc.tables[24].cell(2, 2).text = "若出现稳定 waiting migration、非冗余解释或明确等价边界则进入正文；不要求预设非平凡正结果"
    doc.tables[24].cell(3, 1).text = "检验自然 workload 中 activity→exposure 的同比、非同比、区间依赖或 null"
    doc.tables[24].cell(3, 2).text = "结果稳定、CI/coverage 可解释并符合冻结分析；不以‘必须非显然’作为数据质量门"
    doc.tables[26].cell(2, 2).text = "若存在稳定信息差异或明确等价边界，且 B 能解释；不预设必须误判"
    doc.tables[26].cell(3, 2).text = "允许同比、稳定非同比、区间依赖或 null；主张按结果限定"
    doc.tables[28].cell(4, 1).text = "证明 conventional activity metrics 何时足够、何时不能唯一或可靠推出 exposure；强反转不是必要条件"
    doc.tables[30].cell(2, 1).text = "保留 Q0 方法贡献，按冻结结果收缩 Information Gain；不得看完 Formal 结果后重选 workload"
    doc.tables[32].cell(7, 1).text = "N1/G1 给出传统指标与 exposure 的稳定非冗余差异，或明确、可复现的等价边界；B 解释对应机制"
    set_table(doc.tables[34], [
        ["历史质量项", "历史记录", "v1.4.1 下的资格判断"],
        ["Raw/S/A/B version", "exposedpath-v2", "旧 prototype 身份；不等价于当前方法"],
        ["A conservation", "历史输出曾闭合", "只说明旧分区实现，不能证明 W(s)/terminal 正确"],
        ["GPU/API correlation", "历史报告为 100%", "缺 source manifest，不能证明当前 ownership"],
        ["physical sync UID", "历史输出无重复", "只能用于 Engineering 回归"],
        ["Raw/B interval", "历史输出无一致性错误", "旧 B 定义已变化，必须重新分析"],
        ["局部异常", "4090 p128/b1 曾出现 TRACE_ORDERING_ERROR", "保留为 Q0/回归负例线索"],
    ])

    insert_block_after(find_paragraph(doc, "1.4 ExposedPath 的研究立意与贡献边界"), [
        ("核心问题：在延迟敏感、异步 LLM 推理中，Host 与 GPU 完成了多少活动，与其中多少通过 synchronization/completion dependency 暴露到 request/Prefill/Decode wall-clock，并不是同一个测量对象。ExposedPath 测量后者。", "Decision"),
        ("研究意义：常规 activity 指标可能在某些执行区间足以代表 exposure，也可能出现稳定非同比或区间依赖关系。ExposedPath 不预设答案，而是提供可审计的测量方法，确定两者何时等价、何时失配以及失配由什么同步机制形成。", "Normal"),
        ("Single GPU 自然构成当前基础执行域，因为一次请求已经包含 Host 控制、CUDA 提交、异步 GPU 执行、重叠、同步依赖、hidden progress 与 exposed wait。多 GPU 和在线 serving 会另外引入通信、排队、调度和共享 ownership，需要扩展 S 层语义；它们不是证明当前核心问题成立的前提。", "Normal"),
        ("核心贡献边界限定为单 GPU、请求内部的 Host–device exposure。多 GPU、分布式 serving、并发 ownership、硬件因果归因和通用性能预测均不在当前 claim 内。", "Normal"),
    ])
    insert_block_after(find_paragraph(doc, "2.3 S 层：同步语义归一化"), [
        ("2.3.0 Completion set、W(s) 与时间关系", "Heading 3"),
        ("S 层先依据同步原语的 completion semantics 确定 semantic predecessor set，再用提交证据和 request ownership 构造 W(s)。时间戳用于区分成员在 sync 前已完成、与 sync 重叠或异常晚于 sync；时间重叠本身永远不能生成 dependency。", "Normal"),
    ])
    insert_block_after(find_paragraph(doc, "2.4.2 D 的导航作用与 Exposure Signature"), [
        ("Exposure Signature = {Host path，CUDA submit，CUDA non-submit，device-wait kernel/MemOp/mixed，B hidden/exposed/terminal provenance}。它是冻结 A/B 的机制摘要，不增加新的墙钟预算。", "Normal"),
    ])
    insert_block_after(find_paragraph(doc, "3.3 正式证据主线：Q0→N1/G1→G2"), [
        ("证据链必须按 Correctness→Information Gain→Decision Gain 递进。Q0 不通过时不得解释 N1/G1/G2；N1/G1 没有显示非冗余信息时，G2 不得声称来自 ExposedPath 的决策优势。", "Decision"),
    ])
    insert_block_after(find_paragraph(doc, "3.4 三幅核心图、论文证据链与 Artifact"), [
        ("Figure 1 说明 Raw trace 如何经 S 恢复 W(s)/terminal/validity 并形成 A/B；Figure 2 用 G1 检验 activity 与 exposure 何时同比、非同比或依赖执行区间；Figure 3 用 G2 检验这些新增信息是否影响真实优化响应。三图分别对应测得是否正确、多知道了什么、这些信息是否有用。", "Normal"),
    ])
    insert_block_after(find_paragraph(doc, "4.1 最终可能形成的贡献"), [
        ("以下贡献均为待证据支持的目标，而不是当前已完成事实。最高可防守主张由实际通过的 Gate 决定：Q0 支持正确性，N1/G1 支持信息增量，G2 支持决策增量；任何 null 或失败都必须触发对应 claim 收缩。", "Warning"),
    ])
    insert_block_after(find_paragraph(doc, "4.5 最终停止与降级条件"), [
        ("研究成败由证据等级决定：Q0 失败时停止扩展实验并修正测量语义；Q0 通过但 N1/G1 显示传统指标始终充分时，保留方法正确性与等价边界，重新评估 Information Gain 主张；出现稳定非冗余信息时形成核心实证贡献；只有 held-out G2 进一步增加优化解释或优先级信息时，才声称 Decision Gain。", "Decision"),
    ])

    # 避免相关工作总表只在上一页留下表头和首行，保持章节阅读连续性。
    find_paragraph(doc, "1.2 现有研究脉络与 ExposedPath 的位置").paragraph_format.page_break_before = True

    # 2.6 只是第二章内部小节，无需强制换页；4.6 的短结论块保持在同一页。
    find_paragraph(doc, "2.6 B 层：终止活动生命周期与同步暴露必须分开").paragraph_format.page_break_before = False
    claim_heading = find_paragraph(doc, "4.6 最终研究主句")
    claim_heading.paragraph_format.keep_with_next = True
    paragraphs = doc.paragraphs
    claim_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is claim_heading._p)
    paragraphs[claim_index + 1].paragraph_format.keep_together = True
    paragraphs[claim_index + 1].paragraph_format.keep_with_next = False

    remove_title_rule(doc)
    normalize_revision_marks(doc)
    set_update_fields(doc)
    doc.core_properties.title = "ExposedPath 研究设计"
    doc.core_properties.subject = "当前研究设计主体 v7.1"
    doc.core_properties.comments = "v1.4.1 的背景、立意、方法和执行路线已逐章归并；实验状态仍为 Engineering/Pre-Pilot。"
    doc.save(output)


def patch_protocol(source: Path, output: Path) -> None:
    if sha256(source) != PROTOCOL_SHA256:
        raise RuntimeError("实验协议母版哈希不匹配，拒绝继续")
    shutil.copy2(source, output)
    doc = Document(output)

    replacements = {
        2: "实验协议",
        3: "从 Measurement Contract、Q0 到 Pilot、Protocol Freeze、N1/G1 与 G2 的执行规范",
        4: "依据：《ExposedPath 研究设计》v7.1；方法链 Raw→S→{A,B}→D/Exposure Signature",
        6: "文档职责：本协议承接研究设计中的方法与证据链，规定 Gate、数据角色、候选 WMPC、Q0、N1、G1/G2、Pilot 和冻结规则，不重复维护研究背景。后续实质变更通过文内版本号、修订记录和 Git 差异管理。",
        7: "本协议是当前 Pre-Pilot 执行依据，但不是 Protocol Freeze。它规定 Gate 顺序、数据角色、候选 WMPC、Q0/N1/G1/G2 设计与冻结要求；repeat、overhead、质量门、正式平台和最终矩阵必须由平台资格、OOM/可行域与 Pilot 决定。",
        8: "版本：v2.1　日期：2026 年 9 月 10 日　状态：Pre-Pilot 整合修订版（未冻结 Formal 协议）",
        10: "v1.4.1 整合结论：W(s) 是 synchronization/completion semantics 确定的完整请求内语义前驱集合，包括 sync 入口前已完成成员；valid-empty 不表示前驱已完成；terminal 与 validity 必须基于可验证且唯一的完成证据。",
        24: "Protocol Freeze 前必须由证据确定并冻结：Measurement Contract、Q0 cases/oracle、Canonical Raw 与 analyzer 版本、目标平台、共同可行 WMPC、sentinel、repeat、Pass、quality gates、N1 variants、G2 规则、统计脚本和图表映射。",
        26: "冻结后的 runner/analyzer/registry/oracle、measurement、样本选择或统计规则发生实质变化时，必须生成新协议版本，并使受影响的 Formal 数据失效；Prototype、Engineering 或 Pilot 数据不能通过重命名升级。",
        27: "1.6 从 Measurement Contract 到最终 Protocol Freeze",
        30: "WMPC = Workload×Model×Platform×Computing stack。W01–W22、S1–S6、模型、平台与执行栈在本版中均为候选池，只有通过平台资格、OOM/可行域和 Pilot 后才进入 Protocol Freeze。",
        35: "W01–W22 是待验证的工作负载候选池；G1 Formal 的最终点集、平台与配对数量由共同可行域和 Pilot 统一确定。G2 与可选边界实验优先复用冻结 sentinel。",
        37: "下表保留由预定义切片去重得到的 22 个候选点。Pilot 只允许依据预先声明的显存、稳定性、数据质量和机制覆盖规则整组调整；不得根据显著性、趋势或结果是否‘好看’挑点。",
        52: "C1 是 G2 的优先干预候选，不是本版已经冻结的唯一正式干预。Pilot 必须检查输出一致性、compile/capture 稳定性、graph break/capture coverage、重复波动和 Pass0 收益；不满足条件时删除或收缩 G2，不自动切换跨栈 backend。",
        53: "若需要同栈 fallback，必须在 Formal 数据前明确唯一配置、单变量边界、预测目标和新 Pilot，并发布协议新版本。Llama、vLLM、llama.cpp 等完整栈候选不替代同栈干预。",
        61: "Q0 的核心不是 A 相加等于 NVTX。Q0 必须分别验证：Raw 事实与身份完整；S 的 completion scope、W(s)、terminal 和 validity 符合 CUDA 语义；A/B 对独立 oracle 成立；证据不足、无关重叠和外部 ownership 能正确 fail closed。",
        64: "submitted 的操作判据必须由 Measurement Contract 按可观察事实冻结。活动是否在 sync.start 前完成不影响其语义前驱身份；e.end 与 sync.start 只用于计算 hidden/exposed。若 enqueue 与 sync 的顺序或 ownership 无法证明，输出 AMBIGUOUS/INVALID，不猜成员。",
        66: "per-sync schema 至少包含 sync/request/phase identity、sync_universe_class、completion_scope、registry_rule_id、submission_evidence、dependency_closure_status、wait_set_status、wait_set_activity_ids、terminal evidence/status、validity、ownership、cross_phase_dependency 和原因码。最终字段名由 Measurement Contract 冻结。",
        68: "Validity 至少区分 valid、ambiguous 与 invalid；若使用 VALID_EMPTY，它只能表示没有请求内、受支持且可归属的语义前驱。completed-before-sync 必须保留在 W(s) 并体现 hidden progress。B_NOT_APPLICABLE、B_AMBIGUOUS 与 B_INVALID 分开统计。",
        71: "Q0 通过后只能声明：analyzer 在明确的 observation contract 与同步类型范围内，能够恢复预定义 completion set、W(s)、terminal/validity 和 A/B，并在不可判定时拒绝回答。不能由闭合或相关性推出 Host/CPU/GPU 根因。",
        74: "Q0 用例必须覆盖 Stream、Device/context、Event、completed-before-sync、真正无前驱的 valid-empty、无关跨流重叠、传递依赖、terminal 并列、缺失映射、ownership、dropped records、submission race 和 phase boundary。每个 expected 在运行 analyzer 前写定。",
        87: "N1 只在 Q0 通过后运行，用人为 completion-boundary intervention 检验局部 raw sync 是否足以代表 request-level 增量。G1 的自然逐 Token Token-ready 同步必须保持原行为并独立标识；Token 到文本、网络、前端和外部消费者不进入核心窗口。",
        89: "N1 候选干预限定在 C0=PyTorch eager，并以独立 NVTX/variant/callsite 标识插入同步。具体层位和频率先由 Engineering/Pilot 验证可审计性；干预不得进入 G1 自然 workload，也不得改变输出语义。",
        94: "强信息增量证据的一种可能表现是：干预 callsite 的局部等待增加，而后续自然同步、hidden progress 或 request-level ΔE2E 不按局部 raw sync 一比一变化。该现象不是必需成功条件，必须超过冻结噪声与 profiler overhead 后才能解释。",
        95: "若干预破坏提前提交并产生真实新增 E2E，A/B 能闭合描述新增路径，也可形成有效压力案例；若传统 raw sync 已足够解释，则应报告该场景的信息等价边界。",
        96: "若所有变体中 raw sync 与 request-level 增量近似一一对应，N1 降级为鲁棒性/边界结果，不声称传统指标不可替代。",
        97: "若候选干预没有稳定可观察差异、marker 开销不可控、输出不一致或 stream/callsite 无法唯一审计，则停止扩展干预，不以增加变体追求正结果。",
        99: "5.1 G1：自然工作负载中的信息增量",
        105: "G2 的主预测量、方向判据、held-out 单元和 baseline 必须在 Formal C0/C1 数据前写入 analysis plan。本版保留旧方案作为候选，不宣告预注册完成；Pilot 只验证可行性与噪声。",
        106: "候选 G2_score=A_host_path_ms+A_cuda_api_submit_ms 仅代表待审查的 Host/API 可压缩预算假设，不是 speedup 上界或 measurement definition。若 Measurement Contract/Pilot 发现其不可操作，应在 Formal 前明确删除或版本化修订，不能看完 Formal 结果后调整。",
        107: "候选 baseline 包括 API count、kernel count、total kernel duration、raw sync duration、GPU active ratio 与 launch-to-start；最终集合必须在 Protocol Freeze 中固定，并与 ExposedPath 使用相同配对数据。",
        108: "响应 accounting 候选为 ΔT 与各 A 类变化，并检查 ΣΔA≈ΔT；该闭合用于变化 QA，不证明因果。",
        109: "机制标签只允许由冻结 A/B 与实际干预响应派生；标签集合和判定规则必须在 Formal 前固定。",
        110: "方向、排序、区间或其他 G2 结果指标必须由 Pilot 的可行性与统计设计预先确定；不得把单个 sync 或 kernel 当独立样本。",
        111: "只有 held-out 结果相对冻结 baseline 提供稳定增量，并且 ΔA 与实际 ΔT/主要收缩类别一致时，才能声称 Decision Gain。否则删除预测/决策主张，只保留响应 accounting 或负结果。",
        112: "正式干预优先限定为同栈 eager→compile/graph。POD-Attention、APEX、SPIN、offloading 等只用于讨论机制空间，不自动增加当前矩阵。",
        118: "G3/G4 只在 Q0、N1/G1、G2 主线完成后按论文 claim 选择一个必要适用边界。G5–G7、MoE、量化、llama.cpp 和额外逐 Token 窗口分析均进入后续扩展池，不影响核心进度。",
        125: "5.8 核心证据包与可选适用边界",
        127: "5.9 Pre-Pilot 就绪检查表",
        128: "Measurement Contract 尚未完成：sync universe、required fields、default-stream、completion scope、W(s)、terminal、validity 和原因码仍需版本化审查。",
        129: "新版 analyzer 尚未实现符合 v1.4.1 的 S/A/B；旧 VALID_* 与 accounting 不能计入当前 Gate。",
        130: "Q0 独立 oracle 与真实受控 CUDA trace 尚未完成；合成单元测试不能替代 Q0。",
        131: "W01–W22、S1–S6、模型、平台、repeat 与阈值仍是 Pre-Pilot 候选；必须经平台资格、OOM/可行域和 Pilot 冻结。",
        132: "G2_score、held-out P1 与 baseline 仍是待审查候选，尚未完成 Formal 前预注册。",
        133: "当前先完成 Measurement Contract→Q0 oracle→Canonical Raw→S→A/B；Q0 通过并完成 runner 对齐后才能进入 Engineering Pilot。",
        134: "固定执行顺序：Gate0 Prototype 封存→Gate1 Measurement Contract→Gate2 Q0 oracle→Gate3 Canonical Raw→Gate4 S→Gate5 A/B/D/Signature→Gate6 Q0→Gate7 runner→Gate8 Engineering Pilot→Gate9–12 平台/可行域/Pilot/Freeze→Gate13 N1/G1→Gate14 G2。",
    }
    for index, text in replacements.items():
        set_paragraph(doc, index, text)

    replace_in_headers_and_footers(
        doc,
        "ExposedPath · 实验与分析协议冻结版 v2.0",
        "ExposedPath · 实验与分析协议 v2.1 · Pre-Pilot",
    )

    replace_in_all_text(doc, "研究设计与论文证据框架 v6.0", "ExposedPath 研究设计 v7.1")
    replace_in_all_text(doc, "Pre-Pilot Freeze", "Pre-Pilot 语义修订")

    t0 = doc.tables[0]
    t0.add_row()
    t0.cell(len(t0.rows) - 1, 0).text = "历史 Prototype/Engineering trace"
    t0.cell(len(t0.rows) - 1, 1).text = "旧实现回归、schema 探查和失败线索"
    t0.cell(len(t0.rows) - 1, 2).text = "不能证明当前 W(s)、terminal、A/B 或论文结论"
    t0.cell(len(t0.rows) - 1, 3).text = "不进入 Q0、Pilot 或 Formal"

    set_table(doc.tables[1], [
        ["项目", "Pass0", "Pass1", "Pre-Pilot/冻结要求"],
        ["初始化与 warmup", "使用与 Pass1 相同的候选策略", "含必要 compile/capture；策略与 Pass0 对齐", "编译、capture、缓存建立不进入计时；warmup 数由 Pilot 冻结"],
        ["repeats", "按 Pilot 结果确定", "与同一 WMPC 的 Pass0 配对", "不因显著性追加；repeat 数在 Formal 前冻结"],
        ["主输出", "request/prefill/decode latency、ratio-of-sums throughput", "Raw/S/A/B/D/Signature、coverage、validity、baseline", "Pass1 latency 不替代 Pass0 主性能"],
        ["Profiler overhead", "性能基准", "与 Pass0 同口径比较", "目标、硬门和处理政策由 Pilot 的噪声/成本证据冻结"],
        ["失败重跑", "仅预定义质量失败或整块环境异常", "同左", "按 block 执行，记录 attempt、原因与排除"],
    ])
    gates = [
        ["Gate 0", "封存 Prototype 的提交、trace、环境、输出和限制", "不得升级为当前方法证据"],
        ["Gate 1", "Measurement Contract：phase、sync、W(s)、terminal、validity、A/B/D/Signature", "不得用实现默认值替代未决语义"],
        ["Gate 2", "Q0 独立 oracle 与必需正/负/含糊案例", "不得复用 analyzer 生成 expected"],
        ["Gate 3", "Canonical Raw schema、identity、lineage 与 fail-closed 转换", "下游不得继续各自查询 Nsight 表"],
        ["Gate 4", "实现 S 的 completion set、W(s)、terminal、ownership、validity", "不得把时间重叠当 dependency"],
        ["Gate 5", "实现 A/B，再纯派生 D/Exposure Signature", "不得跨 sync 无规则累加 B"],
        ["Gate 6", "合成与真实 CUDA/Nsight Q0 对独立 oracle", "任一必需 case 失败不得进入 Formal"],
        ["Gate 7", "Runner 的 Token/phase、自然/人为同步、Pass0/1 与跨平台合同", "不得混淆自然 G1 和 N1 干预"],
        ["Gate 8", "Engineering Pilot 打通最小端到端链路", "不得作为科学结论"],
        ["Gate 9", "正式平台 observation stack 与 Q0 资格", "不合格平台不得进入 Formal"],
        ["Gate 10", "OOM、early EOS、稳定性与共同可行域", "OOM 不作为科学意义"],
        ["Gate 11", "Pilot 决定 workload、repeat、overhead、质量门", "不得按结果漂亮程度选择"],
        ["Gate 12", "冻结代码、schema、平台、WMPC、统计与 claim 规则", "冻结后不得静默修改"],
        ["Gate 13", "按冻结协议执行 N1 与自然 G1", "如实报告比例、非比例、区间依赖或 null"],
        ["Gate 14", "执行 held-out G2 并决定最终 claim", "无增量则删除决策优势主张"],
    ]
    set_table(doc.tables[4], [["阶段", "必须完成", "本阶段不得做"], *gates])
    set_table(doc.tables[13], [
        ["术语", "通俗含义", "例子"],
        ["S 的 W(s)", "同步按 completion semantics 必须完成的请求内语义前驱集合；包括 sync 入口前已完成者", "stream sync 的同流提交前缀进入 W(s)，另一流无依赖的重叠活动不进入"],
        ["terminal evidence", "唯一支持同步完成边界的活动或 completion evidence", "只有依赖、identity 和容差都能唯一确定时才给 terminal；否则 ambiguous/invalid"],
    ])
    set_table(doc.tables[15], [
        ["wait_set_status", "terminal/status", "A 行为", "B 行为"],
        ["VALID_NONEMPTY", "唯一 completion evidence 时 VALID", "按 W(s) 与 sync window 的 union 计算；剩余按合同", "B_VALID 或因 terminal 证据不足转 ambiguous"],
        ["VALID_EMPTY", "确无请求内受支持前驱；N/A", "A_device_wait=0；其余按合同", "B_NOT_APPLICABLE"],
        ["AMBIGUOUS", "存在多个合理解释", "不转成零值或已知类别", "B_AMBIGUOUS，保留原因/候选"],
        ["INVALID", "必需事实缺失、冲突或不受支持", "进入 A_unattributed/排除并记录", "B_INVALID"],
    ])
    doc.tables[16].cell(2, 2).text = "owned context 中同步前已 submitted 的请求内语义前驱，包括 sync 入口前已完成成员；ownership 不唯一则 ambiguous/invalid"
    set_table(doc.tables[18], [
        ["Case", "受控结构", "Expected W(s)/terminal", "主要验证"],
        ["Stream sync", "stream1: K1→K2；stream2: K3；同步 stream1", "W={K1,K2}，包括已完成成员；K3 无依赖不进入；terminal 依冻结完成证据", "语义前驱不等于时间重叠或 pending"],
        ["Device/context sync", "owned context 两个 stream 有此前提交活动", "W 包含 scope 内全部请求内语义前驱，包括已完成成员", "跨 stream completion scope 与 ownership"],
        ["Event sync", "K1→record(E)→K2；Host 等待 E", "W 为 event record 所代表的提交前缀，不含 K2", "event record mapping"],
        ["Completed-before-sync", "K1、K2 都在 sync.start 前完成但属于 scope", "K1、K2 仍在 W(s)；暴露交集可为 0；terminal 由冻结规则判断", "hidden progress 不能通过删成员表示"],
        ["Kernel+MemOp mixed", "kernel、memcpy/memset 混合依赖", "成员类型和 terminal evidence 与语义结构一致", "显式 MemOp 支持"],
        ["Unrelated overlap", "另一 stream 的 K3 与 sync 重叠但无 dependency", "K3 不进入 W(s)", "Temporal overlap≠dependency"],
        ["Terminal tie", "多个 completion 候选在冻结容差内不可唯一", "AMBIGUOUS/INVALID；不强选最大时间戳", "terminal fail-closed"],
        ["Missing map / dropped", "移除 event/correlation 或注入 dropped records", "对应 invalid reason；不生成伪 W(s)/terminal", "证据完整性"],
    ])
    set_table(doc.tables[19], [
        ["Case", "受控结构", "Expected", "硬门"],
        ["No predecessor / valid-empty", "同步 scope 内确无请求内、受支持、可归属的前驱", "VALID_EMPTY；W=∅；terminal N/A", "必做"],
        ["Cross-stream wait-event", "stream1: K1→record(E)；stream2: wait(E)→K2；sync stream2", "W 包含 K2 与可观察 K1 前驱", "必做"],
        ["Legacy default stream", "显式构造 legacy default-stream 交互", "按 registry expected closure；mode 不明则 ambiguous/invalid", "必做或明确 unsupported"],
        ["Per-thread default stream", "多 Host 线程构造 PTDS", "与 legacy 预期区分；mode 不明则 ambiguous/invalid", "必做或明确 unsupported"],
        ["Multi-thread ordered", "launch 与 sync 不同线程但提交顺序可证明", "成员与 oracle exact", "必做"],
        ["Submission race", "enqueue 与 sync 顺序不可证明", "AMBIGUOUS；无伪成员", "必做"],
        ["Overlapping Host sync", "两个 Host 线程 sync 区间重叠", "A 用墙钟 union；B 保持 per-sync", "必做"],
        ["Phase-boundary spill", "Prefill-origin activity 在 Decode sync 中完成", "origin/owner/cross_phase 正确；A 只计窗口交集", "必做"],
        ["Invocation bleed", "前一 request activity 进入下一窗口", "标记 ownership/bleed；拒绝强结论", "必做"],
        ["Graph unsupported", "缺 graph node/activity mapping", "GRAPH_MAPPING_UNSUPPORTED；无伪 terminal", "必做"],
        ["Synchronous D2H/implicit", "registry 声明的阻塞 copy/API", "支持则 exact；否则 in-universe unsupported", "纳入时必做"],
        ["Query/poll", "event/stream query 循环", "不因名称自动建立 W(s)", "建议"],
    ])
    doc.tables[34].cell(1, 0).text = "核心可成文证据包"
    doc.tables[34].cell(1, 1).text = "Q0 + G1（N1 作为必要性/边界压力证据）"
    doc.tables[34].cell(1, 2).text = "支持正确性与自然 workload 信息边界；论文等级取决于 G1 的信息增量，不预先承诺投稿层级"
    doc.tables[34].cell(2, 0).text = "高质量核心包"
    doc.tables[34].cell(2, 1).text = "Q0 + N1 + G1 + 有效 G2"
    doc.tables[34].cell(2, 2).text = "Correctness、Information Gain 与 Decision Gain 均有证据"
    doc.tables[34].cell(3, 0).text = "适用边界增强包"
    doc.tables[34].cell(3, 1).text = "高质量核心包 + G3/G4 中必要的一项"
    doc.tables[34].cell(3, 2).text = "增加一个明确的模型或平台适用边界"
    doc.tables[34].cell(4, 0).text = "进一步增强"
    doc.tables[34].cell(4, 1).text = "仅在主线完成后选择 G3/G4 的另一项"
    doc.tables[34].cell(4, 2).text = "不得挤占主证据链，也不自动扩大方法 claim"
    doc.tables[34].cell(5, 0).text = "后续扩展池"
    doc.tables[34].cell(5, 1).text = "G5–G7、MoE、量化、跨引擎、额外平台"
    doc.tables[34].cell(5, 2).text = "当前不执行；另行形成研究问题与协议版本"
    doc.tables[36].cell(3, 1).text = "接受目标、硬门和重采政策由 Pilot 根据噪声、成本与 trace 规模冻结"
    doc.tables[36].cell(5, 1).text = "Pilot 冻结 A_unattributed、supported duration 和 B-valid 门槛；全量报告原因"
    doc.tables[36].cell(6, 1).text = "整数 ns 闭合；API/activity classification coverage 门由 Pilot 冻结，不预设 99%"
    doc.tables[38].cell(2, 1).text = "N1 报告信息等价边界并降级为压力测试；不声称传统指标误判"
    doc.tables[38].cell(3, 1).text = "按冻结结果收缩 Information Gain；不得重新设计或挑选 Formal workload"
    doc.tables[38].cell(7, 1).text = "按 Q0→N1/G1→G2 组织核心证据；仅在需要时增加 G3/G4 边界，G5–G7 延后"

    # 修正实验组总表中的过早正式化与预设正结果。
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                text = text.replace("44 个正式 WMPC", "22×2 个候选端点；Pilot 后冻结实际 Formal 数量")
                text = text.replace("共 44 个正式 WMPC", "形成 22×2 个候选端点；Pilot 后冻结实际 Formal 数量")
                text = text.replace("共 24 个端点", "形成 24 个候选端点；Pilot 后冻结可行子集")
                text = text.replace("发现非显然、可重复的条件化迁移", "允许观察稳定的同比、非同比、区间依赖或 null 关系")
                text = text.replace("证明大部分等待只是迁移", "若后续自然等待相应减少，则支持 waiting migration 解释")
                if text != cell.text:
                    cell.text = text

    insert_block_after(find_paragraph(doc, "第一章　协议范围、共同采集合同与 Protocol Freeze"), [
        ("当前状态：Engineering/Pre-Pilot。Gate0 已通过；Measurement Contract 与 Canonical Raw 尚未完成，Q0、Engineering Pilot、平台资格、Pilot 和 Protocol Freeze 均未通过。历史 trace 只具有 Prototype/Engineering 资格。", "Warning"),
        ("数据角色固定为 Prototype、Engineering、Pilot、Formal 四类。每类数据携带 run role 与 lineage，复制、重命名或重新分析不能提高证据资格。", "Decision"),
    ])
    insert_block_after(find_paragraph(doc, "第三章　方法资格测试（Q0）"), [
        ("Q0 是方法资格门而不是分数。它同时验证 completion semantics、W(s)、terminal、validity、A/B 与 fail-closed；合成 fixture 通过不能替代目标 Nsight observation stack 上的真实受控 CUDA trace。", "Decision"),
    ])
    insert_block_after(find_paragraph(doc, "第四章　同步干预与比较必要性实验（N1）"), [
        ("自然逐 Token 同步与 N1 人为同步必须使用不同 variant/callsite 身份。前者是模型侧 Token-ready completion boundary，属于自然 G1；后者是受控干预。Token 文本化、网络与前端不进入两者的核心测量窗口。", "Decision"),
    ])
    insert_block_after(find_paragraph(doc, "第五章　正式与补充实验组（G1–G7）"), [
        ("当前核心只保留 N1/G1/G2：N1 与 G1承担 Information Gain，G2 承担 Decision Gain。G3/G4 是主线完成后的适用边界候选；G5–G7 统一延后，不作为本轮完成条件。", "Warning"),
    ])
    insert_block_after(find_paragraph(doc, "第六章　Baseline、统计质量、输出 Schema 与停止规则"), [
        ("本章区分测量语义与运行政策：Raw/S/A/B 的定义由 Measurement Contract 冻结；repeat、overhead、coverage、容差和统计门由 Pilot 形成证据后冻结。二者都不能依据 Formal 结果反向调整。", "Decision"),
    ])

    # 避免最后一行结论表单独落在末页；让整个停止/降级映射在新页完整呈现。
    find_paragraph(doc, "6.4 成功、停止、降级与 Claim 映射").paragraph_format.page_break_before = True

    # 横向实验矩阵最后一行较长，禁止只把行尾单字拆到下一页。
    keep_table_row_together(doc.tables[12].rows[-1])

    remove_title_rule(doc)
    normalize_revision_marks(doc)
    set_update_fields(doc)
    doc.core_properties.title = "ExposedPath 实验协议"
    doc.core_properties.subject = "依据 ExposedPath 研究设计 v7.1 的 Pre-Pilot 执行协议"
    doc.core_properties.comments = "尚未 Protocol Freeze；候选 WMPC、平台、阈值与统计规则须经 Pilot。"
    doc.save(output)


def package_inventory(path: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            result[name] = hashlib.sha256(archive.read(name)).hexdigest()
    return result


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit("用法：patch_full_documents.py 研究母版 协议母版 研究输出 协议输出")
    design_source, protocol_source, design_output, protocol_output = map(lambda value: Path(value).resolve(), sys.argv[1:])
    design_output.parent.mkdir(parents=True, exist_ok=True)
    protocol_output.parent.mkdir(parents=True, exist_ok=True)
    patch_design(design_source, design_output)
    patch_protocol(protocol_source, protocol_output)
    manifest = {
        "schema": "exposedpath-document-revision-lineage/1.0",
        "design": {
            "source_filename": design_source.name,
            "source_sha256": sha256(design_source),
            "output": design_output.relative_to(Path.cwd()).as_posix(),
            "output_sha256": sha256(design_output),
            "source_package_parts": len(package_inventory(design_source)),
            "output_package_parts": len(package_inventory(design_output)),
        },
        "protocol": {
            "source_filename": protocol_source.name,
            "source_sha256": sha256(protocol_source),
            "output": protocol_output.relative_to(Path.cwd()).as_posix(),
            "output_sha256": sha256(protocol_output),
            "source_package_parts": len(package_inventory(protocol_source)),
            "output_package_parts": len(package_inventory(protocol_output)),
        },
    }
    manifest_path = Path(__file__).resolve().parent / "revision_lineage_full.json"
    with manifest_path.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(json.dumps(manifest, ensure_ascii=False, indent=2))
        handle.write("\n")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
