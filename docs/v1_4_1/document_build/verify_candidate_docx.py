"""对 ExposedPath 候选 DOCX 做无需 Office 的结构与语义烟雾检查。"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from docx import Document


COMMON_REQUIRED = (
    "Pre-Pilot",
    "不是 Protocol Freeze",
    "Activity Cost",
    "Request-Visible Exposure",
    "W(s)",
    "terminal",
    "validity",
    "Q0",
    "N1",
    "G1",
    "G2",
)


def document_text(doc: Document) -> str:
    chunks = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def verify(path: Path) -> list[str]:
    errors: list[str] = []
    if not path.exists() or path.stat().st_size < 10_000:
        return [f"文件缺失或过小：{path}"]

    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            errors.append(f"OOXML 压缩成员损坏：{bad_member}")
        names = set(archive.namelist())
        for required in ("[Content_Types].xml", "word/document.xml", "word/styles.xml"):
            if required not in names:
                errors.append(f"缺少 OOXML 成员：{required}")

    doc = Document(path)
    text = document_text(doc)
    for phrase in COMMON_REQUIRED:
        if phrase not in text:
            errors.append(f"缺少关键短语：{phrase}")
    for placeholder in ("TODO", "TBD", "FIXME", "待补充", "此处填写"):
        if placeholder in text:
            errors.append(f"残留占位符：{placeholder}")

    if len(doc.paragraphs) < 25:
        errors.append("正文段落数量异常偏少")
    if len(doc.tables) < 1:
        errors.append("表格数量异常偏少")
    for section in doc.sections:
        width_cm = section.page_width.cm
        height_cm = section.page_height.cm
        if abs(width_cm - 21.0) > 0.1 or abs(height_cm - 29.7) > 0.1:
            errors.append(f"页面不是 A4：{width_cm:.2f}×{height_cm:.2f} cm")

    if "研究设计" in path.name:
        for phrase in ("Raw → S → A/B → D / Exposure Signature", "Correctness", "Information Gain", "Decision Gain"):
            if phrase not in text:
                errors.append(f"研究设计缺少：{phrase}")
    if "实验与分析协议" in path.name:
        for phrase in ("W01", "W22", "S1", "S6", "Canonical Raw", "独立 oracle"):
            if phrase not in text:
                errors.append(f"实验协议缺少：{phrase}")

    return errors


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("用法：verify_candidate_docx.py 文件1.docx [文件2.docx ...]")
    failed = False
    for raw_path in sys.argv[1:]:
        path = Path(raw_path).resolve()
        errors = verify(path)
        if errors:
            failed = True
            print(f"FAIL {path}")
            for error in errors:
                print(f"  - {error}")
        else:
            doc = Document(path)
            print(f"PASS {path}：{len(doc.paragraphs)} 个段落，{len(doc.tables)} 个表格")
    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
