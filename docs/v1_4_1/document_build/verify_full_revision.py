"""验证 ExposedPath 两份完整整合修订版 DOCX 的谱系、结构与关键语义。"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document


DESIGN_SHA256 = "0EE058EE1DD2D1FB69EC59D230DBD6B6F03BAEDC319D9C246AEE3ED229B530F1"
PROTOCOL_SHA256 = "1AD97E338F03C951E46BD81B15A5633D46A3B6E6490FEC4B2DFA8840028C0685"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def all_text(doc: Document) -> str:
    values = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    for section in doc.sections:
        for container in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            values.extend(paragraph.text for paragraph in container.paragraphs)
            for table in container.tables:
                for row in table.rows:
                    values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(archive.namelist())
            if name.startswith("word/media/") and not name.endswith("/")
        }


def section_signature(doc: Document) -> list[dict[str, int | None]]:
    result: list[dict[str, int | None]] = []
    for section in doc.sections:
        result.append({
            "width": section.page_width,
            "height": section.page_height,
            "top": section.top_margin,
            "bottom": section.bottom_margin,
            "left": section.left_margin,
            "right": section.right_margin,
            "header": section.header_distance,
            "footer": section.footer_distance,
        })
    return result


def package_checks(path: Path) -> list[str]:
    failures: list[str] = []
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            failures.append(f"ZIP CRC 异常：{bad}")
        names = set(archive.namelist())
        for required in ("[Content_Types].xml", "word/document.xml", "word/styles.xml"):
            if required not in names:
                failures.append(f"缺少 DOCX 必需部件：{required}")
        document_xml = archive.read("word/document.xml").decode("utf-8")
        for marker in ("ins", "del", "moveFrom", "moveTo"):
            if re.search(rf"<w:{marker}(?:[ >])", document_xml):
                failures.append(f"仍含修订标记：<w:{marker}>")
        highlight_values = re.findall(r'<w:highlight w:val="([^"]+)"\s*/>', document_xml)
        unexpected_highlights = sorted(set(highlight_values) - {"none"})
        if unexpected_highlights:
            failures.append(f"仍含可见高亮：{unexpected_highlights}")
        if "word/comments.xml" in names:
            failures.append("仍含批注部件 word/comments.xml")
    return failures


def contains_all(text: str, phrases: list[str]) -> list[str]:
    return [phrase for phrase in phrases if phrase not in text]


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit("用法：verify_full_revision.py 研究母版 协议母版 研究输出 协议输出")

    design_source, protocol_source, design_output, protocol_output = [Path(arg).resolve() for arg in sys.argv[1:]]
    failures: list[str] = []

    if sha256(design_source) != DESIGN_SHA256:
        failures.append("研究设计母版哈希变化")
    if sha256(protocol_source) != PROTOCOL_SHA256:
        failures.append("实验协议母版哈希变化")

    for output in (design_output, protocol_output):
        if not output.exists() or output.stat().st_size == 0:
            failures.append(f"输出不存在或为空：{output}")
        else:
            failures.extend(f"{output.name}: {item}" for item in package_checks(output))

    if failures:
        print(json.dumps({"status": "FAIL", "failures": failures}, ensure_ascii=False, indent=2))
        raise SystemExit(1)

    design_src_doc = Document(design_source)
    protocol_src_doc = Document(protocol_source)
    design_doc = Document(design_output)
    protocol_doc = Document(protocol_output)
    design_text = all_text(design_doc)
    protocol_text = all_text(protocol_doc)

    structural = {
        "design": {
            "paragraphs": len(design_doc.paragraphs),
            "tables": len(design_doc.tables),
            "sections": len(design_doc.sections),
            "media": len(media_hashes(design_output)),
        },
        "protocol": {
            "paragraphs": len(protocol_doc.paragraphs),
            "tables": len(protocol_doc.tables),
            "sections": len(protocol_doc.sections),
            "media": len(media_hashes(protocol_output)),
        },
    }

    if structural["design"]["paragraphs"] <= len(design_src_doc.paragraphs):
        failures.append("研究设计未保留母版并增加逐章整合说明")
    if structural["protocol"]["paragraphs"] <= len(protocol_src_doc.paragraphs):
        failures.append("实验协议未保留母版并增加逐章整合说明")
    if structural["design"]["tables"] != 36 or structural["design"]["sections"] != 1:
        failures.append("研究设计表格数或节数偏离母版")
    if structural["protocol"]["tables"] != 39 or structural["protocol"]["sections"] != 3:
        failures.append("实验协议表格数或节数偏离母版")
    if media_hashes(design_source) != media_hashes(design_output):
        failures.append("研究设计图片集合或图片内容变化")
    if media_hashes(protocol_source) != media_hashes(protocol_output):
        failures.append("实验协议图片集合或图片内容变化")
    if section_signature(design_src_doc) != section_signature(design_doc):
        failures.append("研究设计页面/边距/页眉页脚距离变化")
    if section_signature(protocol_src_doc) != section_signature(protocol_doc):
        failures.append("实验协议横纵节或页面设置变化")

    design_required = [
        "Activity Cost 不等于 Request-Visible Exposure",
        "时延敏感的异步 Host–accelerator 执行系统",
        "传统性能指标与 ExposedPath 的测量对象",
        "传统 trace 和 profiler 主要回答发生了哪些活动",
        "Activity 到 Exposure 之间缺少可验证映射",
        "Raw→S→{A,B}→D/Exposure Signature",
        "时间重叠本身永远不能生成 dependency",
        "VALID_NONEMPTY",
        "VALID_EMPTY",
        "AMBIGUOUS",
        "INVALID",
        "Correctness→Information Gain→Decision Gain",
        "自然逐 Token",
        "Token-ready completion boundary",
        "单 GPU、请求内部",
        "Single GPU 自然构成当前基础执行域",
        "研究不预设必须出现强反转或反常识结果",
        "研究成败由证据等级决定",
        "历史产物已封存为 Prototype/Engineering 证据",
    ]
    protocol_required = [
        "实验与分析协议 v2.1 · Pre-Pilot",
        "依据：《ExposedPath 研究设计》v7.1",
        "不重复维护研究背景",
        "当前状态：Engineering/Pre-Pilot",
        "Measurement Contract 与 Canonical Raw 尚未完成",
        "No predecessor / valid-empty",
        "Cross-stream wait-event",
        "Submission race",
        "Overlapping Host sync",
        "Phase-boundary spill",
        "Q0 是方法资格门而不是分数",
        "自然逐 Token 同步与 N1 人为同步",
        "N1 与 G1承担 Information Gain",
        "G2 承担 Decision Gain",
        "Formal 数据前",
    ]
    for phrase in contains_all(design_text, design_required):
        failures.append(f"研究设计缺少关键语义：{phrase}")
    for phrase in contains_all(protocol_text, protocol_required):
        failures.append(f"实验协议缺少关键语义：{phrase}")

    banned = [
        "44 个正式 WMPC",
        "共 44 个正式 WMPC",
        "证明大部分等待只是迁移",
        "当前已通过 Q0",
        "当前已完成 Protocol Freeze",
        "实验与分析协议冻结版 v2.0",
        "可直接选择“无颜色”去除",
        "TODO",
        "TBD",
    ]
    for label, text in (("研究设计", design_text), ("实验协议", protocol_text)):
        for phrase in banned:
            if phrase in text:
                failures.append(f"{label}仍含过时或占位表述：{phrase}")

    result = {
        "status": "PASS" if not failures else "FAIL",
        "source_hashes": {
            "design": sha256(design_source),
            "protocol": sha256(protocol_source),
        },
        "output_hashes": {
            "design": sha256(design_output),
            "protocol": sha256(protocol_output),
        },
        "structure": structural,
        "failures": failures,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
